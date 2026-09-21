from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
TCC_DIR = ROOT / "docs" / "tcc"
OUT_DIR = TCC_DIR / "output"
REQ_MD = TCC_DIR / "08_requisitos.md"
SCRIPT_MD = TCC_DIR / "02_pitch_video_prototipo.md"
OUT_DOCX = OUT_DIR / "documentacao_requisitos_roteiro_vigilancia_epi.docx"
OUT_REQ_DOCX = OUT_DIR / "documentacao_requisitos_vigilancia_epi.docx"
OUT_SCRIPT_DOCX = OUT_DIR / "roteiro_video_prototipo_vigilancia_epi.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(89, 89, 89)
GRAY_FILL = "F2F4F7"
BLUE_GRAY_FILL = "E8EEF5"
BORDER = "CBD5E1"


def set_cell_text(cell, text, bold=False, size=9.0, color=RGBColor(0, 0, 0), align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    add_inline_runs(p, text, bold_default=bold, size=size, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.append(tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[min(idx, len(widths_dxa) - 1)])
            set_cell_margins(cell)


def set_table_borders(table, color=BORDER, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Pagina ")
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.167


def add_inline_runs(paragraph, text, bold_default=False, size=None, color=None):
    text = text.replace("`", "")
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = bold_default
        value = part
        if part.startswith("**") and part.endswith("**"):
            bold = True
            value = part[2:-2]
        run = paragraph.add_run(value)
        set_run_font(run, size=size, color=color, bold=bold)


def add_body_paragraph(doc, text, style=None, bold=False, italic=False, color=None, after=6, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    add_inline_runs(p, text, bold_default=bold, color=color)
    if italic:
        for r in p.runs:
            r.italic = True
    return p


def _next_numbering_id(numbering, tag_name, attr_name):
    values = []
    for item in numbering.findall(qn(tag_name)):
        raw = item.get(qn(attr_name))
        if raw and raw.isdigit():
            values.append(int(raw))
    return (max(values) + 1) if values else 1


def _decimal_abstract_num_id(doc):
    existing = getattr(doc, "_vigilancia_decimal_abstract_id", None)
    if existing is not None:
        return existing

    numbering = doc.part.numbering_part.element
    abstract_id = _next_numbering_id(numbering, "w:abstractNum", "w:abstractNumId")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))

    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")

    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)

    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal")
    lvl.append(fmt)

    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1.")
    lvl.append(text)

    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    lvl.append(p_pr)

    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        insert_at = list(numbering).index(first_num)
        numbering.insert(insert_at, abstract)
    else:
        numbering.append(abstract)
    doc._vigilancia_decimal_abstract_id = abstract_id
    return abstract_id


def new_numbering_sequence(doc):
    numbering = doc.part.numbering_part.element
    num_id = _next_numbering_id(numbering, "w:num", "w:numId")
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))

    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(_decimal_abstract_num_id(doc)))
    num.append(abstract_ref)

    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    override.append(start_override)
    num.append(override)

    numbering.append(num)
    return num_id


def add_numbered_item(doc, text, index):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    prefix = p.add_run(f"{index}. ")
    set_run_font(prefix, size=11, color=RGBColor(0, 0, 0))
    add_inline_runs(p, text)
    return p


def add_metadata_row(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"{label}: ")
    set_run_font(r, bold=True, size=11, color=RGBColor(0, 0, 0))
    r = p.add_run(value)
    set_run_font(r, size=11, color=RGBColor(0, 0, 0))


def add_callout(doc, text, fill=BLUE_GRAY_FILL):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent_dxa=120)
    set_table_borders(table, color="D7DEE8", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    set_cell_text(cell, text, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def split_table_row(line):
    raw = line.strip().strip("|").split("|")
    return [item.strip() for item in raw]


def is_separator(line):
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells)


def table_widths(headers):
    count = len(headers)
    normalized = [h.lower() for h in headers]
    if count == 2:
        if "parte" in normalized[0] or "perfil" in normalized[0]:
            return [2600, 6760]
        return [2000, 7360]
    if count == 3:
        if "tempo" in normalized[0]:
            return [1600, 3860, 3900]
        if "item" in normalized[0]:
            return [2500, 6860, 0][:3]
        return [1200, 6260, 1900]
    if count == 4:
        return [1050, 5200, 1350, 1760]
    total = 9360
    base = total // count
    widths = [base] * count
    widths[-1] += total - sum(widths)
    return widths


def add_md_table(doc, rows):
    headers = split_table_row(rows[0])
    body = [split_table_row(r) for r in rows[2:]]
    widths = table_widths(headers)
    if len(widths) != len(headers):
        total = 9360
        widths = [total // len(headers)] * len(headers)
        widths[-1] += total - sum(widths)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths, indent_dxa=120)
    set_table_borders(table)

    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, GRAY_FILL)
        align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or len(text) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
        set_cell_text(cell, text, bold=True, size=8.8, color=INK, align=align)

    for source_row in body:
        row = table.add_row()
        set_row_cant_split(row)
        for i in range(len(headers)):
            value = source_row[i] if i < len(source_row) else ""
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or value in {"Alta", "Media", "Baixa", "Implementado", "Parcial"} else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[i], value, size=8.4 if len(headers) >= 4 else 8.8, align=align)
        if len(headers) == 4 and any((source_row[i] if i < len(source_row) else "") == "Implementado" for i in range(len(headers))):
            pass

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    return table


def add_markdown(doc, md_text, skip_first_h1=False):
    lines = md_text.splitlines()
    i = 0
    first_h1_skipped = False
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|") and is_separator(lines[i + 1]):
            table_lines = [stripped, lines[i + 1].strip()]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_md_table(doc, table_lines)
            continue

        if stripped.startswith("#### "):
            p = add_body_paragraph(doc, stripped[5:], after=4)
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = DARK_BLUE
            i += 1
            continue
        if stripped.startswith("### "):
            add_body_paragraph(doc, stripped[4:], style="Heading 3", after=4)
            i += 1
            continue
        if stripped.startswith("## "):
            add_body_paragraph(doc, stripped[3:], style="Heading 2", after=6)
            i += 1
            continue
        if stripped.startswith("# "):
            if skip_first_h1 and not first_h1_skipped:
                first_h1_skipped = True
            else:
                add_body_paragraph(doc, stripped[2:], style="Heading 1", after=8)
            i += 1
            continue
        if stripped.startswith("- "):
            add_body_paragraph(doc, stripped[2:], style="List Bullet", after=4)
            i += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            counter = 1
            while i < len(lines):
                current = lines[i].strip()
                match = re.match(r"^(\d+)\.\s+(.*)$", current)
                if not match:
                    break
                add_numbered_item(doc, match.group(2), counter)
                counter += 1
                i += 1
            continue

        add_body_paragraph(doc, stripped)
        i += 1


def add_cover(doc, subtitle, callout_text):
    section = doc.sections[0]
    header = section.header
    header_p = header.paragraphs[0]
    header_p.text = "Vigilancia EPI | TCC - Desenvolvimento de Sistemas"
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header_p.runs:
        set_run_font(run, size=9, color=MUTED)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.text = ""
    add_page_number(footer_p)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(26)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Vigilancia EPI")
    set_run_font(r, size=27, bold=True, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(subtitle)
    set_run_font(r, size=14, color=MUTED)

    add_metadata_row(doc, "Instituicao", "SENAI Lauro de Freitas - Bahia")
    add_metadata_row(doc, "Curso", "Desenvolvimento de Sistemas")
    add_metadata_row(doc, "Entrega", "Prototipo funcional para TCC")
    add_metadata_row(doc, "Projeto", "Sistema de vigilancia inteligente para uso de EPI")
    add_metadata_row(doc, "Data", "07 de junho de 2026")

    add_callout(
        doc,
        callout_text,
        fill="F4F6F9",
    )

    doc.add_page_break()


def new_doc():
    doc = Document()
    configure_styles(doc)
    return doc


def build_combined():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = new_doc()
    add_cover(
        doc,
        "Documentacao de requisitos e roteiro do video demonstrativo",
        "Este documento organiza os requisitos do prototipo e o roteiro de apresentacao em video. "
        "A proposta do Vigilancia EPI e apoiar a seguranca do trabalho com cameras, modelos de visao computacional e alertas de EPI ausente em tempo real.",
    )

    add_body_paragraph(doc, "Parte 1 - Documentacao de requisitos", style="Heading 1")
    add_markdown(doc, REQ_MD.read_text(encoding="utf-8"), skip_first_h1=True)

    doc.add_page_break()
    add_body_paragraph(doc, "Parte 2 - Roteiro do video demonstrativo", style="Heading 1")
    add_markdown(doc, SCRIPT_MD.read_text(encoding="utf-8"), skip_first_h1=True)

    doc.save(OUT_DOCX)
    return OUT_DOCX


def build_requirements():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = new_doc()
    add_cover(
        doc,
        "Documentacao de requisitos",
        "Este documento apresenta o escopo, as regras de negocio, os requisitos funcionais e nao funcionais, os criterios de aceite e as limitacoes do prototipo Vigilancia EPI.",
    )
    add_markdown(doc, REQ_MD.read_text(encoding="utf-8"), skip_first_h1=False)
    doc.save(OUT_REQ_DOCX)
    return OUT_REQ_DOCX


def build_script():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = new_doc()
    add_cover(
        doc,
        "Roteiro do video demonstrativo do prototipo",
        "Este documento orienta a gravacao do video demonstrativo do Vigilancia EPI, com a ordem das telas, a fala sugerida e os cuidados para apresentar o prototipo de forma natural.",
    )
    add_markdown(doc, SCRIPT_MD.read_text(encoding="utf-8"), skip_first_h1=False)
    doc.save(OUT_SCRIPT_DOCX)
    return OUT_SCRIPT_DOCX


def build():
    paths = [build_requirements(), build_script()]
    for path in paths:
        print(path)


if __name__ == "__main__":
    build()
